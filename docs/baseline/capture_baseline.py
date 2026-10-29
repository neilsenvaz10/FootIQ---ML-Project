"""Non-destructive Milestone 1 capture. Creates new files, never overwrites.

Uses existing packages only. Legacy programs run from fresh temporary copies.
Reproduction: use a NEW --out directory; the default refuses existing evidence.
"""
from __future__ import annotations
import argparse
import ast
from collections import Counter, defaultdict
import csv
from datetime import datetime, timezone
import hashlib
import importlib.metadata as md
import io
import json
import locale
import os
from pathlib import Path
import platform
import re
import shutil
import subprocess
import sys
import tempfile
import time
import zipfile

ROOT = Path(__file__).resolve().parents[2]
parser = argparse.ArgumentParser()
parser.add_argument("--out", type=Path, default=Path(__file__).parent / "capture_2026-10-20")
args = parser.parse_args()
OUT = args.out.resolve()
OUT.mkdir(parents=True, exist_ok=False)
SKIP = {".git", ".venv", "venv", "__pycache__", ".pytest_cache", "node_modules", "provider_cache"}

def sha(path):
    h = hashlib.sha256()
    with path.open("rb") as f:
        for block in iter(lambda: f.read(1024 * 1024), b""):
            h.update(block)
    return h.hexdigest()

def emit(name, data):
    target = OUT / name
    target.parent.mkdir(parents=True, exist_ok=True)
    with target.open("x", encoding="utf-8", newline="") as f:
        if name.endswith(".json"):
            json.dump(data, f, ensure_ascii=False, indent=2, default=str, allow_nan=False)
            f.write("\n")
        else:
            f.write(data)

def originals():
    paths = []
    for base, dirs, names in os.walk(ROOT):
        dirs[:] = [d for d in dirs if d not in SKIP and (Path(base) / d).resolve() != (ROOT / "docs/baseline").resolve()]
        for name in names:
            p = Path(base) / name
            if name == ".env" or (name.startswith(".env.") and name != ".env.example"):
                continue
            paths.append(p)
    return sorted(paths)

def manifest(paths):
    return [{"path": p.relative_to(ROOT).as_posix(), "size_bytes": p.stat().st_size,
             "sha256": sha(p)} for p in paths]

paths = originals()
before = manifest(paths)
emit("original_files_before.json", {"captured_at_utc": datetime.now(timezone.utc).isoformat(), "files": before,
     "exclusions": sorted(SKIP | {"docs/baseline", ".env and non-example variants"})})
emit("original_files.sha256", "".join(f"{r['sha256']}  {r['path']}\n" for r in before))
print(f"Hashed {len(paths)} original project files", flush=True)

provenance = {
    "international_results": {"status": "inherited claim; publisher chain/rights unverified", "evidence": ["data/processed/data_manifest.csv", "outputs/world_cup_data_inventory.md"], "source": "Existing international results archive; inherited manifest cites Kaggle/CC0"},
    "statsbomb": {"status": "local paths verified; remote rights/releases not tested", "source": "StatsBomb Open Data", "url": "https://github.com/statsbomb/open-data", "evidence": ["scripts/download_statsbomb_worldcup.py", "outputs/world_cup_data_inventory.md"]},
    "openfootball": {"status": "local snapshot verified; exact upstream revision unrecorded", "source": "OpenFootball", "url": "https://github.com/openfootball", "evidence": ["data/raw/openfootball/worldcup_2026.json"]},
    "worldbank": {"status": "inherited metadata, rights not newly verified", "source": "World Bank WDI", "evidence": ["data/processed/data_manifest.csv"]},
    "fifa": {"status": "reference HTML snapshot; not a verified clean ranking feed", "source": "FIFA", "evidence": ["data/processed/data_manifest.csv"]},
    "squads": {"status": "public HTML snapshot and heuristic derived values", "source": "Wikipedia World Cup squad tables", "evidence": ["scripts/build_player_impact_v21.py", "data/processed/player_impact_v21_source_audit.csv"]},
    "processed": {"status": "derived/curated; source fields are claims, not independently verified", "source": "Repository scripts and manual inputs", "evidence": ["scripts/", "data/inputs/", "outputs/"]},
}
emit("source_provenance.json", provenance)
NULL = {"", "na", "n/a", "nan", "null", "none"}
date_re = re.compile(r"^\d{4}-\d{2}-\d{2}(?:$|[T ])")

def profile_rows(rows, columns=None):
    columns = list(columns or sorted({str(k) for r in rows if isinstance(r, dict) for k in r}))
    nulls, blanks, dates, invalid = Counter(), Counter(), defaultdict(list), Counter()
    count, malformed = 0, 0
    for r in rows:
        if not isinstance(r, dict):
            malformed += 1
            continue
        count += 1
        for c in columns:
            v = r.get(c)
            st = "" if v is None else str(v).strip()
            blanks[c] += v is None or st == ""
            nulls[c] += v is None or st.lower() in NULL
            if re.search(r"date|timestamp|kickoff|fetched|generated|as_of|valid_from|valid_to", c, re.I) and st:
                try:
                    if date_re.match(st):
                        dt = datetime.fromisoformat(st[:10]).date().isoformat()
                        dates[c].append(dt)
                    else:
                        invalid[c] += 1
                except ValueError:
                    invalid[c] += 1
    return {"row_count": count, "columns": columns,
            "blank_or_null_by_column": dict(blanks), "missing_token_by_column": dict(nulls),
            "date_ranges": {c: {"min": min(v), "max": max(v), "parseable_rows": len(v)} for c, v in dates.items()},
            "non_iso_date_values_by_column": dict(invalid), "non_record_items": malformed}

def json_tables(obj):
    tables = []
    def visit(x, pointer="$"):
        if isinstance(x, list):
            if x and all(isinstance(v, dict) for v in x):
                tables.append({"json_path": pointer, **profile_rows(x)})
            for i, v in enumerate(x):
                if isinstance(v, (dict, list)):
                    visit(v, f"{pointer}[{i}]")
        elif isinstance(x, dict):
            for k, v in x.items():
                if isinstance(v, (dict, list)):
                    visit(v, f"{pointer}.{k}")
    visit(obj)
    return tables

inventory = []
for p in paths:
    rel = p.relative_to(ROOT).as_posix()
    if not rel.startswith(("data/", "outputs/")):
        continue
    item = {"path": rel, "size_bytes": p.stat().st_size, "sha256": sha(p), "format": p.suffix.lower(),
            "row_count": None, "columns": None, "missingness_definition": "blank/null plus case-insensitive NA,N/A,NaN,null,None tokens; tokens may be legitimate text",
            "provenance_key": next((k for k in provenance if k in rel), "processed"),
            "date_policy": "ISO dates in date/time-named fields; date ranges are observed data values, not proven retrieval/availability cutoffs"}
    try:
        if p.suffix.lower() in {".csv", ".tsv"}:
            with p.open(encoding="utf-8-sig", newline="") as f:
                reader = csv.DictReader(f, delimiter="\t" if p.suffix == ".tsv" else ",")
                rows = list(reader)
                item.update(profile_rows(rows, reader.fieldnames))
                item["rows_with_extra_fields"] = sum(None in r for r in rows)
                item["duplicate_headers"] = [k for k, v in Counter(reader.fieldnames or []).items() if v > 1]
        elif p.suffix.lower() == ".json":
            obj = json.loads(p.read_text(encoding="utf-8-sig"))
            item["top_level_type"] = type(obj).__name__
            item["top_level_keys"] = list(obj) if isinstance(obj, dict) else None
            item["tables"] = json_tables(obj)
            if isinstance(obj, list) and all(isinstance(r, dict) for r in obj):
                item.update(profile_rows(obj))
            elif isinstance(obj, list):
                item["top_level_item_count"] = len(obj)
        elif p.suffix.lower() == ".xlsx":
            import openpyxl
            wb = openpyxl.load_workbook(p, read_only=True, data_only=False)
            sheets = []
            for ws in wb:
                vals = list(ws.iter_rows(values_only=True))
                headings = [str(v) if v is not None else f"unnamed_{i}" for i, v in enumerate(vals[0])] if vals else []
                sheets.append({"name": ws.title, "header_assumption": "first worksheet row; display/report sheets may not be rectangular data",
                               "max_row": ws.max_row, "max_column": ws.max_column,
                               **profile_rows([dict(zip(headings, r)) for r in vals[1:]], headings)})
            wb.close()
            item["sheets"] = sheets
        else:
            item["profiling_note"] = "Non-tabular/report/binary asset; row/column/missing-value counts not applicable"
        item["profile_status"] = "ok"
    except Exception as exc:
        item["profile_status"] = "error"
        item["profile_error"] = f"{type(exc).__name__}: {exc}"
    inventory.append(item)
emit("dataset_inventory.json", {"files": inventory, "note": "JSON nested table counts overlap; do not sum tables as distinct matches. Spreadsheet missingness assumes first-row headers."})
with (OUT / "dataset_inventory.csv").open("x", encoding="utf-8", newline="") as f:
    w = csv.DictWriter(f, fieldnames=["path", "size_bytes", "sha256", "format", "row_count", "column_count", "missing_cells", "date_ranges", "provenance_key", "profile_status"])
    w.writeheader()
    for r in inventory:
        w.writerow({**{k: r.get(k) for k in ("path", "size_bytes", "sha256", "format", "row_count", "provenance_key", "profile_status")},
                    "column_count": len(r["columns"]) if r["columns"] else None,
                    "missing_cells": sum(r.get("missing_token_by_column", {}).values()),
                    "date_ranges": json.dumps(r.get("date_ranges", {}), sort_keys=True)})
print(f"Profiled {len(inventory)} data/report assets", flush=True)

scripts = []
for p in paths:
    if p.suffix != ".py":
        continue
    source = p.read_text(encoding="utf-8-sig")
    tree = ast.parse(source)
    imports, functions, refs, urls, io_calls = [], [], [], [], []
    for node in ast.walk(tree):
        if isinstance(node, ast.Import):
            imports.extend(a.name for a in node.names)
        elif isinstance(node, ast.ImportFrom):
            imports.append(node.module or "")
        elif isinstance(node, (ast.FunctionDef, ast.AsyncFunctionDef)):
            functions.append({"name": node.name, "line": node.lineno})
        elif isinstance(node, ast.Constant) and isinstance(node.value, str):
            if re.search(r"\.(csv|json|pkl|xlsx|html|md)$", node.value):
                refs.append({"literal": node.value, "line": node.lineno})
            if node.value.startswith("http"):
                urls.append(node.value)
        elif isinstance(node, ast.Call):
            name = ast.unparse(node.func)
            if any(token in name.lower() for token in ["write", "read", "open", "dump", "load", "fit", "seed"]):
                io_calls.append({"call": ast.unparse(node)[:500], "line": node.lineno})
    scripts.append({"path": p.relative_to(ROOT).as_posix(), "imports": sorted(set(imports)), "functions": functions,
                    "file_literals": refs, "urls": sorted(set(urls)), "io_and_fit_calls": io_calls})
emit("script_dependency_map.json", {"analysis": "Static AST references; literals are not proof of read/write direction. Runtime traces establish exercised file dependencies.", "scripts": scripts})

packages = sorted([{"name": d.metadata.get("Name", "unknown"), "version": d.version} for d in md.distributions()], key=lambda x: x["name"].lower())
env = {"python_executable": sys.executable, "python_version": sys.version, "implementation": platform.python_implementation(),
       "platform": platform.platform(), "machine": platform.machine(), "architecture": platform.architecture(),
       "prefix": sys.prefix, "base_prefix": sys.base_prefix, "default_encoding": sys.getdefaultencoding(),
       "filesystem_encoding": sys.getfilesystemencoding(), "preferred_encoding": locale.getpreferredencoding(False),
       "utf8_mode": sys.flags.utf8_mode, "packages": packages,
       "secrets": "No environment variable dump, credentials or provider responses captured",
       "requirements_sha256": sha(ROOT / "requirements-share.txt"),
       "portability": "Version inventory, not an offline wheelhouse; hashes of distributions and cross-platform installation not verified"}
emit("environment.json", env)
emit("requirements-observed.txt", "".join(f"{p['name']}=={p['version']}\n" for p in packages))
pip_cmd = [sys.executable, "-B", "-X", "utf8", "-m", "pip", "freeze", "--all"]
pip = subprocess.run(pip_cmd, capture_output=True, text=True, encoding="utf-8", errors="replace")
emit("pip-freeze-command.json", {"command": pip_cmd, "exit_code": pip.returncode, "stdout": pip.stdout, "stderr": pip.stderr,
     "note": "The existing venv may lack pip; importlib.metadata inventory remains authoritative for installed versions."})

# Exact source snapshot retains existing data/output bytes, excludes secrets/runtime.
archive = OUT / "reference_snapshot.zip"
with zipfile.ZipFile(archive, "x", compression=zipfile.ZIP_DEFLATED, compresslevel=6) as z:
    for p in paths:
        z.write(p, p.relative_to(ROOT).as_posix())
emit("snapshot_manifest.json", {"archive": archive.name, "sha256": sha(archive), "size_bytes": archive.stat().st_size,
     "files": before, "usage": "Local reproducibility copy; inherited third-party rights are unresolved. Not cleared for redistribution."})

work = Path(tempfile.mkdtemp(prefix="isolated_execution_", dir=OUT))
emit("execution_workspace.json", {"path": str(work), "policy": "New temporary copies retained; no deletion. No network or legacy algorithm modifications."})
commands = []
actions = ["build_live_predictions_v27.py", "build_confidence_gate_v28.py", "simulate_live_tournament_v213.py", "validate_worldcup_full_v214.py"]
for number in (1, 2):
    runroot = work / f"run{number}"
    runroot.mkdir()
    with zipfile.ZipFile(archive) as z:
        z.extractall(runroot)
    if number == 1:
        actions_run = ["pytest", *actions]
    else:
        actions_run = actions[:3]
    for action in actions_run:
        trace = runroot / f"audit_trace_{action}.json"
        cmd = [sys.executable, "-B", "-X", "utf8", str(Path(__file__).parent / "isolated_runner.py"), str(runroot), action, str(trace)]
        started = time.perf_counter()
        proc = subprocess.run(cmd, capture_output=True, text=True, encoding="utf-8", errors="replace", env={**os.environ, "PYTHONDONTWRITEBYTECODE": "1", "PYTHONUTF8": "1", "PYTEST_DISABLE_PLUGIN_AUTOLOAD": "1"})
        name = f"logs/run{number}_{action}"
        emit(name + ".stdout.txt", proc.stdout)
        emit(name + ".stderr.txt", proc.stderr)
        record = {"run": number, "action": action, "command": cmd, "exit_code": proc.returncode, "duration_seconds": round(time.perf_counter()-started, 3)}
        commands.append(record)
        if trace.exists():
            emit(f"traces/run{number}_{action}.json", json.loads(trace.read_text(encoding="utf-8")))
        print(json.dumps(record), flush=True)
    # Capture outputs touched by the runs, not every duplicated input.
    generated = sorted({r for t in runroot.glob("audit_trace_*.json") for r in json.loads(t.read_text(encoding="utf-8"))["written_paths"]})
    for rel in generated:
        if (runroot / rel).is_file():
            target = OUT / f"references/run{number}" / rel
            target.parent.mkdir(parents=True, exist_ok=True)
            with target.open("xb") as f:
                f.write((runroot / rel).read_bytes())
    protected = [r for r in before if r["path"].startswith(("scripts/", "tests/", "data/raw/", "data/model_artifacts/"))]
    emit(f"run{number}_protected_integrity.json", {"checked": len(protected), "changed": [r["path"] for r in protected if sha(runroot / r["path"]) != r["sha256"]]})
emit("commands.json", commands)

refs = ["data/processed/v27_live_model_features.csv", "data/processed/v27_live_probabilities.csv", "data/processed/v27_live_public_picks.csv", "data/processed/v28_live_public_picks.csv", "data/processed/live_tournament_v213_tournament_odds.csv"]
comparisons = []
for rel in refs:
    a, b = work / "run1" / rel, work / "run2" / rel
    original = ROOT / rel
    comparisons.append({"path": rel, "run1_sha256": sha(a), "run2_sha256": sha(b), "exact_repeat_match": sha(a)==sha(b),
                        "original_sha256": sha(original) if original.exists() else None,
                        "original_byte_match": sha(a)==sha(original) if original.exists() else None})
emit("determinism_comparison.json", {"numeric_tolerance": 1e-12, "policy": "Exact bytes compared for stable CSVs; timestamps in simulator config/reports are intentionally not deterministic.", "comparisons": comparisons})

# Inspect local trusted model and rebuild historical features without training/writing.
meta_cmd = [sys.executable, "-B", "-X", "utf8", "-c", "import pickle,json,sys; from pathlib import Path; sys.path.insert(0,sys.argv[1]); import model_bakeoff_v26 as v; a=pickle.load(open(sys.argv[2],'rb')); d=v.build_feature_frame(); splits=v.split_frame(d); m={k:x for k,x in a.items() if k!='models'}; m['models']={k:{'type':type(x).__module__+'.'+type(x).__name__,'parameters':x.get_params()} for k,x in a['models'].items()}; m['feature_columns_code']=v.FEATURE_COLUMNS; m['rebuilt_feature_rows']=len(d); m['split_counts']=[len(x) for x in splits]; m['split_date_ranges']=[{'min':str(x.date.min()),'max':str(x.date.max())} for x in splits]; print(json.dumps(m,default=str))", str(work / "run1/scripts"), str(work / "run1/data/model_artifacts/v27/live_match_model.pkl")]
meta = subprocess.run(meta_cmd, capture_output=True, text=True, encoding="utf-8", errors="replace", env={**os.environ, "PYTHONDONTWRITEBYTECODE": "1"})
emit("artifact_inspection_command.json", {"command": meta_cmd, "exit_code": meta.returncode, "stderr": meta.stderr})
if meta.returncode == 0:
    emit("artifact_metadata.json", json.loads(meta.stdout))
else:
    emit("artifact_inspection_failure.txt", meta.stdout)

after = manifest(paths)
changes = [{"path": a["path"], "before": a["sha256"], "after": b["sha256"]} for a,b in zip(before,after) if a!=b]
new_originals = sorted(set(p.relative_to(ROOT).as_posix() for p in originals()) - {r["path"] for r in before})
emit("original_files_after.json", {"files": after})
emit("integrity_verification.json", {"original_files_checked": len(before), "changed": changes, "new_files_outside_baseline": new_originals,
     "all_originals_unchanged": not changes and not new_originals, "note": "Runtime/caches/git/secrets excluded; audited project files checked byte-for-byte."})
emit("capture_summary.json", {"original_files": len(before), "total_original_bytes": sum(r["size_bytes"] for r in before),
     "dataset_assets": len(inventory), "formats": dict(Counter(r["format"] for r in inventory)),
     "profile_errors": [{"path":r["path"],"error":r.get("profile_error")} for r in inventory if r["profile_status"]!="ok"],
     "snapshot_bytes": archive.stat().st_size, "commands_failed": [r for r in commands if r["exit_code"]],
     "originals_unchanged": not changes and not new_originals, "archive_sha256": sha(archive)})
print("CAPTURE COMPLETE " + str(OUT), flush=True)
